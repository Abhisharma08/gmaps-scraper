"""Write results to CSV, XLSX, or JSON based on the output file extension."""

import csv
import json
import os
from typing import List

from .models import FIELDNAMES, Business


def write(businesses: List[Business], path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    directory = os.path.dirname(os.path.abspath(path))
    if directory:
        os.makedirs(directory, exist_ok=True)

    if ext == ".json":
        _write_json(businesses, path)
    elif ext in (".xlsx", ".xls"):
        _write_xlsx(businesses, path)
    else:
        _write_csv(businesses, path)
    return path


def _write_csv(businesses: List[Business], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDNAMES)
        writer.writeheader()
        for biz in businesses:
            writer.writerow(biz.to_row())


def _write_json(businesses: List[Business], path: str) -> None:
    from dataclasses import asdict

    with open(path, "w", encoding="utf-8") as fh:
        json.dump([asdict(b) for b in businesses], fh, indent=2, ensure_ascii=False)


def _write_xlsx(businesses: List[Business], path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise SystemExit("XLSX output needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Businesses"
    ws.append(FIELDNAMES)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for biz in businesses:
        row = biz.to_row()
        ws.append([row[name] for name in FIELDNAMES])

    for i, name in enumerate(FIELDNAMES, 1):
        longest = max(
            [len(name)] + [len(str(b.to_row()[name])) for b in businesses] or [len(name)]
        )
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(
            longest + 2, 60
        )

    ws.freeze_panes = "A2"
    wb.save(path)
